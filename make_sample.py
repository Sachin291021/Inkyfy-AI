"""Creates sample.docx for testing: python make_sample.py"""
from docx import Document
d = Document()
d.add_heading("The Water Cycle", 0)
d.add_heading("Introduction", 1)
d.add_paragraph("Water is constantly moving between the oceans, the atmosphere and the land. This continuous journey is called the water cycle, and it keeps every living thing on Earth supplied with fresh water.")
d.add_heading("Main Stages", 1)
for s in ["Evaporation: heat from the sun turns liquid water into invisible vapour.", "Condensation: rising vapour cools and forms tiny droplets, which gather as clouds.", "Precipitation: droplets grow heavy and fall as rain, snow or hail.", "Collection: water flows into rivers, lakes and oceans, and the cycle begins again."]:
    d.add_paragraph(s)
d.add_heading("Conclusion", 1)
d.add_paragraph("Without the water cycle there would be no rain to grow crops, no rivers to drink from, and no clouds to shade us. " * 4)
d.save("sample.docx")
